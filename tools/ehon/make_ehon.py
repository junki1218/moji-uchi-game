"""えほん（Google スライド用の pptx）を組む（2026-10-22 作り直し。最初は 2026-10-05 に一時フォルダで作った）。

    python tools/ehon/make_ehon.py tools/ehon/gakkun_01.json

- 絵: 画像生成で作った pptx（1スライドに絵1枚）から取り出す。または画像ファイルを直接指定
- 文: JSON に書く（1ページ＝絵1枚＋右の余白の文）
- 出力: 16:9 の pptx。絵は左に高さいっぱい、右の余白に文。文の「が」は黒地に白い字（アプリの見本と同じ）
- できた pptx は G: ドライブにコピーし、Google スライドで開いて「ファイル → Google スライドとして保存」で変換する
  （Drive のコネクタでは pptx を Google スライドに変換できないため）

JSON の書き方は tools/ehon/README.md
"""

from __future__ import annotations

import io
import json
import re
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

W, H = Inches(13.333), Inches(7.5)
BG = RGBColor(0xFF, 0xF7, 0xEC)
INK = RGBColor(0x22, 0x22, 0x22)
SOUND = RGBColor(0xE0, 0x7A, 0x1F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Klee One"  # Google スライドで使える教科書体ふうのフォント

# 文の種類 → (大きさ pt, 太字, 色)。Google スライドでは字が少し大きく出るので、1行は6字くらいまでにする
STYLE = {
    "title": (34, True, INK),  # 題名
    "main": (32, True, INK),  # 大事な文（「いぬが……」「いぬが はしる。」）
    "nar": (22, False, INK),  # 地の文
    "snd": (24, True, SOUND),  # 音（「びゅーん！」）
    "gap": (12, False, INK),  # 空き行
}


def set_font(run, size: int, bold: bool, color: RGBColor) -> None:
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, FONT
    f.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):  # 日本語の字にも同じフォントを使わせる
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def highlight_black(run) -> None:
    """「が」を黒地にする（白い字と組み合わせて、アプリの黒丸の「が」に寄せる）"""
    rpr = run._r.get_or_add_rPr()
    hl = etree.SubElement(rpr, qn("a:highlight"))
    etree.SubElement(hl, qn("a:srgbClr")).set("val", "222222")
    fill = rpr.find(qn("a:solidFill"))
    if fill is not None:  # スキーマの順番: solidFill のあとに highlight
        rpr.remove(hl)
        fill.addnext(hl)


def images_from_pptx(path: Path) -> dict[int, bytes]:
    """pptx の各スライドの1枚目の絵 → {スライド番号(1から): 画像データ}"""
    out: dict[int, bytes] = {}
    for n, slide in enumerate(Presentation(str(path)).slides, 1):
        for sh in slide.shapes:
            if sh.shape_type == 13:  # PICTURE
                out[n] = sh.image.blob
                break
    return out


def to_jpeg(data: bytes, crop_ratio: float | None) -> tuple[io.BytesIO, tuple[int, int]]:
    """JPEG にする。crop_ratio（横÷縦）を渡すと、横を真ん中で切ってその比にする"""
    im = Image.open(io.BytesIO(data)).convert("RGB")
    if crop_ratio:
        w, h = im.size
        cw = min(w, round(h * crop_ratio))
        x = (w - cw) // 2
        im = im.crop((x, 0, x + cw, h))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=90)
    buf.seek(0)
    return buf, im.size


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit("使い方: python tools/ehon/make_ehon.py <えほんの JSON>")
    spec_path = Path(sys.argv[1]).resolve()
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    base = spec_path.parent

    def resolve(p: str) -> Path:
        q = Path(p)
        return q if q.is_absolute() else (base / q).resolve()

    slides_src = images_from_pptx(resolve(spec["images_pptx"])) if spec.get("images_pptx") else {}

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    for n, page in enumerate(spec["pages"], 1):
        if "image" in page:
            path = resolve(page["image"])
            if not path.is_file():
                sys.exit(f"{n}ページ目: 絵が見つからない {path}")
            data = path.read_bytes()
        else:
            data = slides_src.get(page["slide"])
            if data is None:
                sys.exit(f"{n}ページ目: 絵の pptx に {page['slide']} 枚目の絵がない")
        buf, (iw, ih) = to_jpeg(data, page.get("crop_ratio"))

        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = BG
        pw = int(H * iw / ih)  # 絵は高さいっぱい
        s.shapes.add_picture(buf, 0, 0, pw, H)

        pad = Inches(0.15)
        tb = s.shapes.add_textbox(pw + pad, Inches(0.3), W - pw - pad * 2, H - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.05)
        for k, (text, kind) in enumerate(page["lines"]):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            size, bold, color = STYLE[kind]
            if not text:
                r = p.add_run()
                r.text = " "
                set_font(r, size, False, INK)
                continue
            for part in re.split(r"(\[が\])", text):
                if not part:
                    continue
                r = p.add_run()
                if part == "[が]":
                    r.text = "が"
                    set_font(r, size, True, WHITE)
                    highlight_black(r)
                else:
                    r.text = part
                    set_font(r, size, bold, color)

    out = resolve(spec["output"])
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"{out}  {len(prs.slides)}ページ  {out.stat().st_size // 1024}KB")


if __name__ == "__main__":
    main()
